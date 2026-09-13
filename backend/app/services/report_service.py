from typing import Dict, Any, List, Optional
from datetime import datetime
import csv
import io
from uuid import uuid4

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.models.database import SensorReading, FailurePrediction
from app.services.predictive_service import (
    PredictiveMaintenanceService,
    EQUIPMENT_LIMITS,
    SYNTHETIC_EQUIPMENT,
)
from app.services.sensor_gateway import TELEMETRY_METRICS

service = PredictiveMaintenanceService()

RISK_FILLS = {
    "CRITICAL": "FFCDD2",
    "HIGH": "FFE0B2",
    "MEDIUM": "FFF9C4",
    "LOW": "C8E6C9",
    "NORMAL": "C8E6C9",
}


def _sensor_readings(equipment_id: str, db, limit: int = 100) -> List[Dict[str, Any]]:
    rows = (
        db.query(SensorReading)
        .filter(SensorReading.equipment_id == equipment_id)
        .order_by(SensorReading.timestamp.asc())
        .limit(limit)
        .all()
    )
    return [
        {
            "timestamp": r.timestamp.isoformat() if r.timestamp else None,
            "vibration": r.vibration,
            "temperature": r.temperature,
            "pressure": r.pressure,
            "current": r.current,
            "noise_level": r.noise_level,
            "operating_hours": r.operating_hours,
            "maintenance_age_days": r.maintenance_age_days,
            "remaining_days": r.remaining_days,
            "is_synthetic": r.is_synthetic,
        }
        for r in rows
    ]


async def build_equipment_report(equipment_id: str, db) -> Dict[str, Any]:
    readings = _sensor_readings(equipment_id, db)
    if not readings:
        raise ValueError(f"No sensor history available for '{equipment_id}'")

    latest = db.query(SensorReading).filter(
        SensorReading.equipment_id == equipment_id
    ).order_by(SensorReading.timestamp.desc()).first()

    prediction = await service.predict_equipment_failure(equipment_id, readings)

    return {
        "equipment_id": equipment_id,
        "equipment_name": latest.equipment_name if latest else equipment_id,
        "generated_at": datetime.utcnow().isoformat(),
        "data_points": len(readings),
        "prediction": prediction,
        "readings": readings[-15:],
    }


def _summary_rows(report: Dict[str, Any]) -> List[List[str]]:
    p = report.get("prediction", {})
    assessment = p.get("assessment", {})
    return [
        ["Equipment ID", report.get("equipment_id", "")],
        ["Equipment", report.get("equipment_name", "")],
        ["Data points", str(report.get("data_points", 0))],
        ["Risk level", assessment.get("risk_level", "—")],
        ["Predicted failure (days)", str(p.get("predicted_failure_days", "—"))],
        ["Confidence", f"{round((p.get('confidence') or 0) * 100, 1)}%"],
        ["Anomalies detected", str(p.get("anomalies_detected", 0))],
        ["Recommended action", p.get("recommended_action", "—")],
        ["Model", p.get("model", "—")],
    ]


# ── PDF ────────────────────────────────────────────────────────────
def render_maintenance_pdf(report: Dict[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom", parent=styles["Title"], fontSize=18, spaceAfter=4,
        textColor=colors.HexColor("#0f172a"),
    )
    header_style = ParagraphStyle(
        "HeaderCustom", parent=styles["Heading2"], fontSize=12,
        textColor=colors.HexColor("#006064"), spaceBefore=14, spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "MetaCustom", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748b"),
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            rightMargin=0.7 * inch, leftMargin=0.7 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    story = []

    story.append(Paragraph("MechMind AI - Equipment Maintenance Report", title_style))
    story.append(Paragraph(f"Generated {report['generated_at']}", meta_style))
    story.append(Spacer(1, 6))

    story.append(Paragraph("Equipment Summary", header_style))
    summary = Table([["Field", "Value"]] + _summary_rows(report), colWidths=[2.1 * inch, 4.7 * inch])
    summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d9488")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
    ]))
    story.append(summary)

    story.append(Paragraph("Latest Sensor Readings", header_style))
    metric_rows = [["Metric", "Value", "Normal Limit", "Critical Limit", "Status"]]
    latest = report["readings"][-1] if report["readings"] else {}
    eq_type = SYNTHETIC_EQUIPMENT.get(report.get("equipment_id"), {}).get(
        "equipment_type", "centrifugal pump"
    )
    eq_type_base = EQUIPMENT_LIMITS.get(
        eq_type, EQUIPMENT_LIMITS["centrifugal pump"]
    )
    for metric in TELEMETRY_METRICS:
        value = latest.get(metric)
        if value is None:
            continue
        limits = _metric_limits(metric, eq_type_base)
        status = "CRITICAL" if value >= limits["critical"] else (
            "ABOVE LIMIT" if value >= limits["limit"] else "OK"
        )
        metric_rows.append([
            metric.replace("_", " ").title(),
            f"{value}",
            str(limits["limit"]),
            str(limits["critical"]),
            status,
        ])
    table = Table(metric_rows, colWidths=[1.7 * inch, 1.0 * inch, 1.3 * inch, 1.3 * inch, 1.5 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d9488")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
    ]))
    story.append(table)

    story.append(Paragraph("Reading History (last 15)", header_style))
    history_rows = [["Timestamp", "Vib", "Temp", "Press", "Curr", "Noise"]]
    for r in report["readings"]:
        ts = (r.get("timestamp") or "")[:19].replace("T", " ")
        history_rows.append([
            ts,
            _fmt(r.get("vibration")),
            _fmt(r.get("temperature")),
            _fmt(r.get("pressure")),
            _fmt(r.get("current")),
            _fmt(r.get("noise_level")),
        ])
    hist = Table(history_rows, colWidths=[2.0 * inch, 0.8 * inch, 0.8 * inch, 0.8 * inch, 0.8 * inch, 0.8 * inch])
    hist.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    story.append(hist)

    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "This report is auto-generated by MechMind AI and is for maintenance planning purposes only. "
        "Predictions are statistical estimates based on available sensor data.",
        ParagraphStyle("Footer", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#94a3b8")),
    ))

    doc.build(story)
    return buf.getvalue()


def _metric_limits(metric: str, base: Dict[str, Dict[str, float]]) -> Dict[str, float]:
    if metric in base:
        return base[metric]
    defaults = {
        "vibration": (0.3, 0.45),
        "temperature": (80.0, 95.0),
        "pressure": (0.0, 1500.0),
        "current": (0.0, 150.0),
        "noise_level": (0.0, 95.0),
    }
    limit, critical = defaults.get(metric, (0.0, 0.0))
    return {"limit": limit, "critical": critical}


def _fmt(v: Optional[float]) -> str:
    return "—" if v is None else f"{round(v, 2)}"


# ── Excel ──────────────────────────────────────────────────────────
def _style_worksheet(ws, headers: List[str], rows: List[List[Any]], widths: List[float]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D9488")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w


def generate_equipment_status_xlsx(db) -> bytes:
    predictions = (
        db.query(FailurePrediction)
        .order_by(FailurePrediction.created_at.desc())
        .all()
    )
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Equipment Status"
    headers = ["Equipment ID", "Equipment", "Risk Level", "RUL (days)",
               "Confidence", "Anomalies", "Recommended Action", "Model",
               "Created At"]
    rows = [
        [
            p.equipment_id,
            p.equipment_name,
            p.risk_level,
            p.predicted_days,
            p.confidence,
            p.anomalies_detected,
            p.recommended_action,
            p.model,
            p.created_at.isoformat() if p.created_at else "",
        ]
        for p in predictions
    ]
    _style_worksheet(ws, headers, rows, [22, 24, 12, 12, 12, 12, 40, 12, 22])

    for idx in range(2, len(rows) + 2):
        risk = ws.cell(row=idx, column=3).value
        fill = RISK_FILLS.get(risk)
        if fill and risk:
            ws.cell(row=idx, column=3).fill = PatternFill("solid", fgColor=fill)

    wb.save(buf)
    return buf.getvalue()


def generate_predictions_xlsx(db) -> bytes:
    return generate_equipment_status_xlsx(db)


# ── CSV ────────────────────────────────────────────────────────────
def _csv_response(data: List[List[str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(data)
    return buf.getvalue().encode("utf-8-sig")


def generate_equipment_status_csv(db) -> bytes:
    predictions = (
        db.query(FailurePrediction)
        .order_by(FailurePrediction.created_at.desc())
        .all()
    )
    rows = [
        ["Equipment ID", "Equipment", "Risk Level", "RUL (days)", "Confidence",
         "Anomalies", "Recommended Action", "Model", "Created At"],
        *[
            [
                p.equipment_id,
                p.equipment_name or "",
                p.risk_level or "",
                p.predicted_days if p.predicted_days is not None else "",
                p.confidence or "",
                p.anomalies_detected or "",
                p.recommended_action or "",
                p.model or "",
                p.created_at.isoformat() if p.created_at else "",
            ]
            for p in predictions
        ],
    ]
    return _csv_response(rows)


def generate_maintenance_csv(equipment_id: str, db) -> bytes:
    readings = _sensor_readings(equipment_id, db)
    if not readings:
        raise ValueError(f"No sensor history available for '{equipment_id}'")

    rows = [
        ["timestamp", "vibration", "temperature", "pressure", "current",
         "noise_level", "operating_hours", "maintenance_age_days", "remaining_days"],
        *[
            [
                r.get("timestamp") or "",
                r.get("vibration") or "",
                r.get("temperature") or "",
                r.get("pressure") or "",
                r.get("current") or "",
                r.get("noise_level") or "",
                r.get("operating_hours") or "",
                r.get("maintenance_age_days") or "",
                r.get("remaining_days") if r.get("remaining_days") is not None else "",
            ]
            for r in readings
        ],
    ]
    return _csv_response(rows)